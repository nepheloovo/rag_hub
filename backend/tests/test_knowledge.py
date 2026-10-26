"""知识库 / 文档导入管道自测（Celery eager 模式，伪向量降级）"""
import io

from app.celery_app import celery_app


def setup_module():
    # 测试中任务同步执行，无需独立 worker
    celery_app.conf.task_always_eager = True


def teardown_module():
    celery_app.conf.task_always_eager = False


def _create_kb(client, name: str = "pytest知识库") -> str:
    resp = client.post("/api/knowledge-bases", json={"name": name, "description": "自测"})
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def _upload(client, kb_id: str, filename: str, content: bytes, mime: str):
    return client.post(
        f"/api/knowledge-bases/{kb_id}/documents",
        files={"file": (filename, io.BytesIO(content), mime)},
    )


def _refetch_doc(client, kb_id: str, doc_id: str) -> dict:
    """上传接口返回创建瞬间的状态，需重新查询拿解析结果"""
    docs = client.get(f"/api/knowledge-bases/{kb_id}/documents").json()
    return next(d for d in docs if d["id"] == doc_id)


def test_kb_crud(client):
    kb_id = _create_kb(client)
    # 列表包含新库
    ids = [k["id"] for k in client.get("/api/knowledge-bases").json()]
    assert kb_id in ids
    # 更新
    resp = client.put(f"/api/knowledge-bases/{kb_id}", json={"description": "改过"})
    assert resp.json()["description"] == "改过"
    # 删除
    assert client.delete(f"/api/knowledge-bases/{kb_id}").status_code == 204
    assert kb_id not in [k["id"] for k in client.get("/api/knowledge-bases").json()]


def test_upload_markdown_pipeline(client):
    kb_id = _create_kb(client)
    md = "# 标题A\n\n## 小节\n\n这是用于测试的知识内容，包含检索目标语句。\n\n## 附录\n\n另一段内容。"
    resp = _upload(client, kb_id, "test.md", md.encode("utf-8"), "text/markdown")
    assert resp.status_code == 201, resp.text
    assert resp.json()["status"] == "pending"  # 返回创建瞬间状态
    doc = _refetch_doc(client, kb_id, resp.json()["id"])
    assert doc["status"] == "done"  # eager 模式下同步完成
    assert doc["chunk_count"] >= 2

    # 分块预览
    chunks = client.get(f"/api/documents/{doc['id']}/chunks").json()
    assert len(chunks) == doc["chunk_count"]
    assert any("检索目标语句" in c["content"] for c in chunks)
    assert all(c["source"] == "test.md" for c in chunks)

    # 删除文档
    assert client.delete(f"/api/documents/{doc['id']}").status_code == 204
    assert client.get(f"/api/documents/{doc['id']}/chunks").status_code == 404


def test_upload_unsupported_type(client):
    kb_id = _create_kb(client)
    resp = _upload(client, kb_id, "evil.exe", b"x", "application/octet-stream")
    assert resp.status_code == 415


def test_upload_image_without_ocr_fails_gracefully(client):
    """未配置 OCR 服务商时，图片导入应失败并给出可读原因"""
    kb_id = _create_kb(client)
    # 1x1 PNG
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
        "0000000d49444154789c626001000000ffff03000006000557bfabd40000000049454e44ae426082"
    )
    resp = _upload(client, kb_id, "pic.png", png, "image/png")
    assert resp.status_code == 201
    doc = _refetch_doc(client, kb_id, resp.json()["id"])
    assert doc["status"] == "failed"
    assert "OCR" in doc["error"]


def test_upload_docx_and_table(client):
    kb_id = _create_kb(client)
    from docx import Document

    buf = io.BytesIO()
    d = Document()
    d.add_heading("制度", level=1)
    d.add_paragraph("包含表格的制度文档。")
    d.save(buf)

    resp = _upload(client, kb_id, "制度.docx", buf.getvalue(),
                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert resp.status_code == 201
    doc = _refetch_doc(client, kb_id, resp.json()["id"])
    assert doc["status"] == "done"
    assert doc["chunk_count"] >= 1

    # CSV
    resp = _upload(client, kb_id, "名单.csv", "姓名,城市\n张三,上海\n".encode("utf-8"), "text/csv")
    assert resp.status_code == 201
    doc = _refetch_doc(client, kb_id, resp.json()["id"])
    assert doc["status"] == "done"
